"""Builds the pediatrician review pack from the app's own sources.

    python docs/pediatrician-review/make_review_docx.py

Writes, next to this script:
    palna-medical-review.md    (Markdown version)
    palna-medical-review.docx  (Word version: A4, tables, page numbers)

Both come from one content model, so they always match. The medical DATA is
read straight from the app (see SOURCES below); the hand-written text
(letter, per-country sources, open questions, explanations) is in
review_content.py next to this script. Where the pack explains a rule in
words (growth check, fever warning, RSV / flu logic), the script first checks
that the rule in the source still looks the same, and stops with a message
if it changed, so the explanation can't silently go stale.

Needs python-docx (pip install --user python-docx).
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True  # keep this folder free of __pycache__

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]  # D:\projects\baby-app
sys.path.insert(0, str(HERE))
import review_content as C  # noqa: E402

SOURCES = {
    "migrations": ROOT / "supabase" / "migrations",
    "optional": ROOT / "app" / "lib" / "core" / "optional_vaccines.dart",
    "emergency": ROOT / "app" / "lib" / "core" / "emergency.dart",
    "arb": ROOT / "app" / "lib" / "l10n" / "app_en.arb",
    "growth_check": ROOT / "app" / "lib" / "features" / "growth" / "growth_check.dart",
    "weight_velocity": ROOT / "app" / "lib" / "core" / "weight_velocity.dart",
    "who_growth": ROOT / "app" / "lib" / "core" / "who_growth.dart",
    "who_velocity_json": ROOT / "app" / "assets" / "who_weight_velocity.json",
    "symptom_sheet": ROOT / "app" / "lib" / "features" / "health" / "symptom_sheet.dart",
    "models": ROOT / "app" / "lib" / "data" / "models.dart",
    "summary_pdf": ROOT / "app" / "lib" / "features" / "health" / "summary_pdf.dart",
    "report_stats": ROOT / "app" / "lib" / "features" / "reports" / "report_stats.dart",
    "sleep_predictions": ROOT / "app" / "lib" / "features" / "home" / "sleep_predictions.dart",
    "foods": ROOT / "app" / "assets" / "foods.json",
    "milestones": ROOT / "app" / "assets" / "milestones.json",
}

OUT_MD = HERE / "palna-medical-review.md"
OUT_DOCX = HERE / "palna-medical-review.docx"

REVIEW_HEADERS = ["Reviewer\n\u2713 / \u2717 / ?", "Change to / comment"]


class SourceChanged(Exception):
    pass


def read(key: str) -> str:
    return SOURCES[key].read_text(encoding="utf-8")


def need(pattern: str, text: str, what: str, flags=0) -> re.Match:
    m = re.search(pattern, text, flags)
    if not m:
        raise SourceChanged(f"Could not find {what}. The source changed shape: update make_review_docx.py "
                            f"and the matching explanation in review_content.py.")
    return m


def expect(key: str, *snippets: str, what: str) -> None:
    """Stops if a rule we explain in words no longer appears in the source."""
    text = re.sub(r"\s+", " ", read(key))
    for s in snippets:
        if re.sub(r"\s+", " ", s) not in text:
            raise SourceChanged(f"{what}: expected to find {s!r} in {SOURCES[key].name}. The rule changed: "
                                f"update the explanation in make_review_docx.py / review_content.py.")


# ---------------------------------------------------------------------------
# SQL: vaccine schedule rows
# ---------------------------------------------------------------------------

def sql_statements(sql: str) -> list[str]:
    """Splits SQL into statements, dropping -- comments, respecting quotes."""
    out, cur, i, n, quote = [], [], 0, len(sql), False
    while i < n:
        ch = sql[i]
        if quote:
            cur.append(ch)
            if ch == "'":
                if i + 1 < n and sql[i + 1] == "'":
                    cur.append("'")
                    i += 1
                else:
                    quote = False
        elif ch == "'":
            quote = True
            cur.append(ch)
        elif ch == "-" and sql.startswith("--", i):
            while i < n and sql[i] != "\n":
                i += 1
            continue
        elif ch == ";":
            out.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
        i += 1
    if "".join(cur).strip():
        out.append("".join(cur).strip())
    return out


def sql_values(text: str) -> list[list]:
    """Parses "(a, 'b', null), (...)" into lists of Python values."""
    rows, row, tok, i, depth, quote = [], [], [], 0, 0, False
    n = len(text)

    def flush():
        raw = "".join(tok).strip()
        tok.clear()
        if raw.startswith("'"):
            row.append(raw[1:-1].replace("''", "'"))
        elif raw.lower() == "null":
            row.append(None)
        elif raw.lower() in ("true", "false"):
            row.append(raw.lower() == "true")
        elif re.fullmatch(r"-?\d+", raw):
            row.append(int(raw))
        else:
            raise SourceChanged(f"Unexpected SQL value {raw!r} in the vaccine schedule.")

    while i < n:
        ch = text[i]
        if quote:
            tok.append(ch)
            if ch == "'":
                if i + 1 < n and text[i + 1] == "'":
                    tok.append("'")
                    i += 1
                else:
                    quote = False
        elif ch == "'":
            quote = True
            tok.append(ch)
        elif ch == "(":
            depth += 1
            if depth == 1:
                row, tok = [], []
            else:
                tok.append(ch)
        elif ch == ")":
            depth -= 1
            if depth == 0:
                flush()
                rows.append(row)
            else:
                tok.append(ch)
        elif ch == "," and depth == 1:
            flush()
        elif depth >= 1:
            tok.append(ch)
        i += 1
    return rows


def sql_conditions(where: str) -> dict:
    conds = {}
    for part in re.split(r"\s+and\s+", where.strip(), flags=re.I):
        m = re.fullmatch(r"(country_code|code)\s*=\s*'([^']*)'", part.strip(), re.I)
        if m:
            conds[m.group(1).lower()] = {m.group(2)}
            continue
        m = re.fullmatch(r"(country_code|code)\s+in\s*\((.*)\)", part.strip(), re.I)
        if m:
            conds[m.group(1).lower()] = set(re.findall(r"'([^']*)'", m.group(2)))
            continue
        raise SourceChanged(f"Can't read the condition {part!r} on the vaccine schedule.")
    return conds


def matches(row: dict, conds: dict) -> bool:
    return all(row[k] in v for k, v in conds.items())


def load_schedules() -> dict[str, list[dict]]:
    rows: list[dict] = []
    for path in sorted(SOURCES["migrations"].glob("*.sql")):
        for st in sql_statements(path.read_text(encoding="utf-8")):
            low = " ".join(st.lower().split())
            if low.startswith("insert into public.vaccine_schedule"):
                m = need(r"insert\s+into\s+public\.vaccine_schedule\s*\(([^)]*)\)\s*values\s*(.*)$", st,
                         "the schedule insert", re.I | re.S)
                cols = [c.strip().lower() for c in m.group(1).split(",")]
                for vals in sql_values(m.group(2)):
                    r = dict(zip(cols, vals))
                    r.setdefault("verified", False)
                    rows.append(r)
            elif low.startswith("update public.vaccine_schedule"):
                m = need(r"update\s+public\.vaccine_schedule\s+set\s+(.*?)(?:\s+where\s+(.*))?$", st,
                         "the schedule update", re.I | re.S)
                sets = {}
                for a in re.split(r",(?=(?:[^']*'[^']*')*[^']*$)", m.group(1)):
                    k, v = a.split("=", 1)
                    sets[k.strip().lower()] = sql_values("(" + v.strip() + ")")[0][0]
                conds = sql_conditions(m.group(2)) if m.group(2) else {}
                for r in rows:
                    if matches(r, conds):
                        r.update(sets)
            elif low.startswith("delete from public.vaccine_schedule"):
                m = re.search(r"where\s+(.*)$", st, re.I | re.S)
                conds = sql_conditions(m.group(1)) if m else {}
                rows = [r for r in rows if not matches(r, conds)]
    by_country: dict[str, list[dict]] = {}
    for r in rows:
        by_country.setdefault(r["country_code"], []).append(r)
    for v in by_country.values():
        v.sort(key=lambda r: r["sort_order"])
    return by_country


def fmt_age(value: int, unit: str) -> str:
    if unit == "day":
        return "At birth" if value == 0 else f"{value} days"
    if unit == "week":
        return f"{value} weeks"
    if value < 24:
        return f"{value} month" + ("s" if value != 1 else "")
    y, mo = divmod(value, 12)
    if mo == 0:
        return f"{y} years ({value} months)"
    if mo == 6:
        return f"{y}\u00bd years ({value} months)"
    return f"{y} years {mo} months ({value} months)"


# ---------------------------------------------------------------------------
# Dart: RSV / flu rules, emergency numbers, thresholds
# ---------------------------------------------------------------------------

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def md_text(month: int, day: int) -> str:
    return f"{day} {MONTHS[month - 1]}"


def load_optional():
    src = read("optional")
    rsv_block = need(r"const rsvRules = <String, RsvRule>\{(.*?)\n\};", src, "the RSV rules", re.S).group(1)
    flu_block = need(r"const fluRules = <String, FluRule>\{(.*?)\n\};", src, "the flu rules", re.S).group(1)
    md = r"\(month:\s*(\d+),\s*day:\s*(\d+)\)"
    rsv = {}
    for m in re.finditer(r"'([A-Z]{2})':\s*\(\s*birthFrom:\s*" + md + r",\s*birthTo:\s*" + md +
                         r",\s*catchUpOn:\s*(null|" + md + r"),\s*catchUpUnderMonths:\s*(null|\d+),?\s*\)",
                         rsv_block):
        g = m.groups()
        rsv[g[0]] = {
            "from": (int(g[1]), int(g[2])), "to": (int(g[3]), int(g[4])),
            "catch": None if g[5] == "null" else (int(g[6]), int(g[7])),
            "under": None if g[8] == "null" else int(g[8]),
        }
    flu = {}
    for m in re.finditer(r"'([A-Z]{2})':\s*\(fromMonths:\s*(\d+),\s*untilMonths:\s*(null|\d+),\s*campaignStart:\s*"
                         + md + r",\s*twoDosesUnderYears:\s*(null|\d+),\s*nasal:\s*(true|false)\)", flu_block):
        g = m.groups()
        flu[g[0]] = {"from": int(g[1]), "until": None if g[2] == "null" else int(g[2]),
                     "start": (int(g[3]), int(g[4])), "two": None if g[5] == "null" else int(g[5]),
                     "nasal": g[6] == "true"}
    if len(rsv) != rsv_block.count("birthFrom") or len(flu) != flu_block.count("fromMonths"):
        raise SourceChanged("Some RSV or flu rules could not be read: update load_optional().")
    # The "re-check every summer" notes at the top of the file.
    notes, cur = [], None
    for line in src.splitlines():
        if line.startswith("///   * "):
            cur = line[len("///   * "):].strip()
            notes.append(cur)
        elif cur is not None and line.startswith("///     "):
            notes[-1] += " " + line[3:].strip()
        else:
            cur = None
    # Logic described in words in the pack.
    expect("optional", "if (!on.isBefore(_addMonths(birth, under))) return null;",
           "final end = _addMonths(start, 6);", "final due = oldEnough.isAfter(start) ? oldEnough : start;",
           "firstFlu: !hadFluBefore && twoDoses != null && due.isBefore(_addMonths(birth, twoDoses * 12))",
           what="RSV / flu reminder logic")
    return rsv, flu, notes


def load_emergency():
    src = read("emergency")
    block = need(r"const _numbers = <String, List<EmergencyNumber>>\{(.*?)\n\};", src, "the emergency numbers",
                 re.S).group(1)
    numbers = {}
    for m in re.finditer(r"'([A-Z]{2})':\s*\[(.*?)\]", block, re.S):
        numbers[m.group(1)] = re.findall(r"\(number:\s*'([^']+)',\s*service:\s*EmergencyService\.(\w+)\)",
                                         m.group(2))
    if len(numbers) != len(re.findall(r"'[A-Z]{2}':", block)) or not all(numbers.values()):
        raise SourceChanged("Some emergency numbers could not be read: update load_emergency().")
    eu = set(re.findall(r"'([A-Z]{2})'", need(r"const _eu112 = \{(.*?)\};", src, "the EU 112 list",
                                              re.S).group(1)))
    pk_primary = need(r"if \(country == 'PK'\) return '([^']+)';", src, "the Pakistan call-now number").group(1)
    heat = set(re.findall(r"'([A-Z]{2})'", need(r"const heatAdviceCountries = \{([^}]*)\}", src,
                                                "the hot-weather countries").group(1)))
    return numbers, eu, pk_primary, heat


SERVICE_LABEL_KEYS = {
    "all": "numAll", "ambulance": "numAmbulance", "rescue1122": "numRescue", "edhi": "numEdhi",
    "chhipa": "numChhipa", "police": "numPolice", "fire": "numFire", "nhs111": "numNhs111",
    "poison": "numPoison", "nurseLine": "numNurseLine", "unified": "numUnified",
    "outOfHoursDoctor": "numOutOfHoursDoctor", "healthAdvice": "numHealthAdvice",
}


def load_thresholds():
    t = {}
    sheet = read("symptom_sheet")
    m = need(r"c >= ([\d.]+) && babyAge\([^)]*\)\.months < (\d+)", sheet, "the young-baby fever rule")
    t["fever_c"], t["fever_months"] = float(m.group(1)), int(m.group(2))
    gc = read("growth_check")
    t["birthweight_days"] = int(need(r"days > (\d+) && latest\.value \* 1000 < baby\.birthWeightG", gc,
                                     "the birth-weight rule").group(1))
    m = need(r"gap < (\d+) \|\| gap > (\d+)", gc, "the curve-crossing window")
    t["cross_min"], t["cross_max"] = int(m.group(1)), int(m.group(2))
    t["cross_lines"] = int(need(r"crossed <= -(\d+)", gc, "the curve-crossing count").group(1))
    expect("growth_check",
           "flags.add(GrowthFlag(metric, GrowthFlagKind.belowRange, GrowthConcern.doctor));",
           "metric == GrowthMetric.head ? GrowthConcern.doctor : GrowthConcern.mention",
           "flags.add(GrowthFlag(metric, GrowthFlagKind.fallingAcross, GrowthConcern.doctor));",
           "GrowthFlagKind.belowBirthWeight, GrowthConcern.doctor",
           "GrowthFlagKind.slowGain, GrowthConcern.mention",
           "p.measurement != null && p.months <= WhoGrowth.maxMonths",
           what="growth check levels")
    who = read("who_growth")
    t["curves"] = [(lbl, float(z)) for lbl, z in re.findall(r"\('(\d+\w\w)', (-?[\d.]+)\)", who)]
    t["who_max_months"] = int(need(r"static const maxMonths = (\d+);", who, "the WHO age limit").group(1))
    wv = read("weight_velocity")
    t["gain_min_days"] = int(need(r"static const minDays = (\d+);", wv, "the weight-gain minimum gap").group(1))
    t["gain_max_days"] = int(need(r"static const maxDays = (\d+);", wv, "the weight-gain maximum gap").group(1))
    expect("weight_velocity", "midAge > 12 * WhoGrowth.daysPerMonth",
           "gainPerMonthG < p5 ? GainBand.slow : (gainPerMonthG > p95 ? GainBand.fast : GainBand.usual)",
           what="weight-gain bands")
    rs = read("report_stats")
    sleep_block = need(r"typicalSleepHours\(int ageDays\) => switch \(ageDays\) \{(.*?)\};", rs,
                       "the typical sleep ranges", re.S).group(1)
    t["sleep"] = [(int(a), int(b), int(c)) for a, b, c in
                  re.findall(r"< (\d+) => \(from: (\d+), to: (\d+)\)", sleep_block)]
    sp = read("sleep_predictions")
    awake_block = need(r"typicalAwakeMinutes\(int ageMonths\) => switch \(ageMonths\) \{(.*?)\};", sp,
                       "the awake windows", re.S).group(1)
    t["awake"] = re.findall(r"^\s*(.+?) => \(min: (\d+), max: (\d+)\)", awake_block, re.M)
    expect("summary_pdf", "Not a medical record or a diagnosis.", what="PDF summary footer")
    return t


def load_symptoms() -> list[str]:
    block = need(r"const symptomCodes = \[(.*?)\];", read("models"), "the symptom list", re.S).group(1)
    return re.findall(r"'([a-z_]+)'", block)


# ---------------------------------------------------------------------------
# Wording helpers
# ---------------------------------------------------------------------------

def resolve_icu(text: str) -> str:
    """Replaces {x, plural, ... other{...}} with its "other" branch."""
    while True:
        m = re.search(r"\{(\w+), plural,", text)
        if not m:
            return text
        depth, i = 0, m.start()
        while i < len(text):
            if text[i] == "{":
                depth += 1
            elif text[i] == "}":
                depth -= 1
                if depth == 0:
                    break
            i += 1
        body = text[m.end():i]
        om = re.search(r"other\{(.*)\}\s*$", body, re.S)
        text = text[:m.start()] + (om.group(1) if om else body) + text[i + 1:]


def fill(text: str) -> str:
    text = resolve_icu(text)
    return re.sub(r"\{(\w+)\}", lambda m: "[" + C.PLACEHOLDERS.get(m.group(1), m.group(1)) + "]", text)


def split_points(text: str) -> list[str]:
    return [ln.strip().lstrip("\u2022").strip() for ln in text.split("\n") if ln.strip()]


# ---------------------------------------------------------------------------
# Content model
# ---------------------------------------------------------------------------

class Doc:
    def __init__(self):
        self.blocks: list[tuple] = []
        self.question_count = 0
        self.table_count = 0

    def add(self, *block):
        self.blocks.append(block)

    def table(self, headers, rows, widths, title=None):
        """[widths] (cm) are for the data columns; the two reviewer columns fill the rest of the 17 cm."""
        if len(widths) != len(headers) or sum(widths) > 17.0 - 1.8 - 1.5:
            raise ValueError(f"Table {title!r}: bad column widths {widths}")
        self.table_count += 1
        self.add("table", headers + REVIEW_HEADERS, [list(r) + ["", ""] for r in rows],
                 list(widths) + [1.8, round(17.0 - sum(widths) - 1.8, 2)], title)

    def questions(self, qs):
        self.question_count += len(qs)
        self.add("h3", "Open questions")
        self.add("questions", qs)

    def sources(self, items):
        self.add("h3", "Sources")
        self.add("bullets", items)


def build() -> Doc:
    schedules = load_schedules()
    rsv, flu, optional_notes = load_optional()
    numbers, eu112, pk_primary, heat_countries = load_emergency()
    th = load_thresholds()
    symptoms = load_symptoms()
    arb = json.loads(read("arb"))
    foods = json.loads(read("foods"))
    ms = json.loads(read("milestones"))
    velocity = json.loads(read("who_velocity_json"))

    def s(key):
        if key not in arb:
            raise SourceChanged(f"The wording {key!r} is no longer in the app: update review_content.py.")
        return fill(arb[key])

    missing = [c for c in C.COUNTRY_ORDER if c not in schedules]
    extra = [c for c in schedules if c not in C.COUNTRY_ORDER]
    if missing or extra:
        raise SourceChanged(f"Countries changed (missing {missing}, new {extra}): update review_content.py.")

    d = Doc()

    # Cover -----------------------------------------------------------------
    d.add("title", C.TITLE)
    d.add("subtitle", "Vaccine schedules, RSV and flu reminders, medical wording, emergency numbers, solids "
                      "guide and milestones in the Palna baby-tracking app")
    d.add("p", f"Version: {C.VERSION_DATE}")
    d.add("signature", ["Prepared for", "Reviewed by", "Qualifications / registration no.", "Signature", "Date"])
    d.add("pagebreak")

    # Letter ------------------------------------------------------------------
    d.add("h1", "A letter to the reviewing doctor")
    for para in C.LETTER:
        d.add("p", para)
    d.add("pagebreak")

    # Top questions -----------------------------------------------------------
    d.add("h1", "Questions we most need answered")
    d.add("p", "If you only have a little time, please answer these first. Each one is repeated, with more "
               "detail, in its section; you can answer in either place.")
    d.add("toplist", C.TOP_QUESTIONS)
    d.add("pagebreak")

    # 1. Vaccine schedules ----------------------------------------------------
    total_doses = sum(len(v) for v in schedules.values())
    d.add("h1", "1. Vaccine schedules (23 countries)")
    for para in C.SCHEDULE_INTRO:
        d.add("p", para)
    d.add("p", f"In total: {total_doses} doses in {len(schedules)} countries. Countries in this section: "
               + ", ".join(f"1.{i + 1} {C.COUNTRY_NAMES[c]}" for i, c in enumerate(C.COUNTRY_ORDER)) + ".")
    for i, code in enumerate(C.COUNTRY_ORDER):
        info = C.COUNTRIES[code]
        rows = schedules[code]
        d.add("pagebreak")
        d.add("h2", f"1.{i + 1} {C.COUNTRY_NAMES[code]}")
        d.add("p", info["note"])
        d.table(["#", "Vaccine", "Dose", "Due age", "Verified"],
                [[f"{code}-{n + 1}", r["name"], r["dose_label"] or "\u2014",
                  fmt_age(r["due_age_value"], r["due_age_unit"]), "Yes" if r["verified"] else "No"]
                 for n, r in enumerate(rows)],
                [1.5, 4.9, 2.0, 2.5, 1.6],
                title=f"{C.COUNTRY_NAMES[code]}: routine schedule ({len(rows)} doses)")
        d.sources(info["sources"])
        d.questions(info["questions"])
    d.add("pagebreak")
    d.add("h2", "1.24 All countries")
    d.sources(C.SCHEDULE_SOURCES)
    d.questions(C.SCHEDULE_QUESTIONS)
    d.add("pagebreak")

    # 2. RSV and flu -----------------------------------------------------------
    d.add("h1", "2. Optional RSV and flu reminders")
    for para in C.OPTIONAL_INTRO:
        d.add("p", para)
    d.add("h2", "2.1 What parents see")
    wording = [("optionalVaccinesTitle", "Heading of the optional list"),
               ("optionalVaccinesHelp", "Under the heading"),
               ("rsvName", "Name of the RSV item"), ("rsvNote", "Note under the RSV item"),
               ("fluName", "Name of the flu item"), ("fluNote", "Note under the flu item"),
               ("fluNoteFirst", "Added the first time (see rule above)"),
               ("fluNoteNasal", "Added in the United Kingdom and Ireland"),
               ("optionalReminderTitle", "Reminder notification, title"),
               ("optionalReminderBody", "Reminder notification, text"),
               ("optionalVaccinesHidden", "When the family hides the reminders")]
    d.table(["#", "Where it appears", "Exact text shown"],
            [[f"O{n + 1}", where, s(k)] for n, (k, where) in enumerate(wording)],
            [1.0, 3.4, 7.4], title="Optional reminder wording")

    d.add("h2", "2.2 RSV protection: which babies get a reminder")
    rows = []
    for code in C.COUNTRY_ORDER:
        if code not in rsv:
            continue
        r = rsv[code]
        window = f"{md_text(*r['from'])} \u2013 {md_text(*r['to'])}"
        catch = ("No catch-up reminder" if r["catch"] is None else
                 f"Reminder on {md_text(*r['catch'])} if under {r['under']} months that day")
        rows.append([C.COUNTRY_NAMES[code], window, catch, C.RSV_BASIS.get(code, "")])
    d.table(["Country", "Born in this window: reminder at birth", "Born before the window", "Programme rule"],
            rows, [2.3, 2.4, 3.0, 4.6], title="RSV reminder rules")

    d.add("h2", "2.3 Flu vaccine: which children get a reminder")
    rows = []
    for code in C.COUNTRY_ORDER:
        if code not in flu:
            continue
        f = flu[code]
        ages = f"From {f['from']} months" + (f" to under {f['until']} months" if f["until"] else ", no upper age")
        rows.append([C.COUNTRY_NAMES[code], ages, md_text(*f["start"]),
                     f"Under {f['two']} years" if f["two"] else "Not shown (1 dose)",
                     "Yes" if f["nasal"] else "No"])
    d.table(["Country", "Ages", "Season reminder from", "Two-dose note (first time)", "Nasal-spray note"],
            rows, [2.4, 3.0, 2.0, 2.6, 1.6], title="Flu reminder rules")

    d.add("h2", "2.4 All 23 countries at a glance")
    d.add("p", "What each country\u2019s programme offers, from our research, and whether the app shows an "
               "optional reminder.")
    d.table(["Country", "RSV: national programme", "RSV in app", "Flu for young children", "Flu in app"],
            [[C.COUNTRY_NAMES[c], C.PROGRAMME_OVERVIEW[c][0], "Yes" if c in rsv else "No",
              C.PROGRAMME_OVERVIEW[c][1], "Yes" if c in flu else "No"] for c in C.COUNTRY_ORDER],
            [2.2, 4.3, 1.4, 3.9, 1.4], title="RSV and flu by country")
    if optional_notes:
        d.add("h3", "Notes kept with these rules (to re-check every summer)")
        d.add("bullets", optional_notes)
    d.sources(C.OPTIONAL_SOURCES)
    d.questions(C.OPTIONAL_QUESTIONS)
    d.add("pagebreak")

    # 3. Medical wording -------------------------------------------------------
    d.add("h1", "3. Medical wording shown to parents")
    for para in C.WORDING_INTRO:
        d.add("p", para)
    counter = [0]

    def wording_table(items, title):
        rows = []
        for key, where, split in items:
            parts = split_points(s(key)) if split else [s(key)]
            for j, part in enumerate(parts):
                counter[0] += 1
                rows.append([f"W{counter[0]}", where if j == 0 else "\u2033 (continued)", part])
        d.table(["#", "Where it appears", "Exact text shown"], rows, [1.0, 3.4, 7.4], title=title)

    for heading, intro, items in C.WORDING_GROUPS:
        d.add("h2", heading)
        if intro:
            d.add("p", intro)
        if heading.startswith("3.1"):
            d.add("p", f"Rule: the warning appears when a temperature of {th['fever_c']:g} \u00b0C "
                       f"({th['fever_c'] * 9 / 5 + 32:.1f} \u00b0F) or more is entered for a baby under "
                       f"{th['fever_months']} completed months at the time of the reading. Temperatures can be "
                       f"entered in \u00b0C or \u00b0F; the app does not ask where the temperature was taken.")
        wording_table(items, heading[4:])
        if heading.startswith("3.4"):
            sym_rows = [[f"W{counter[0] + 1}", "Symptoms a parent can log (with an optional temperature and "
                                               "severity)",
                         ", ".join(s("symptom" + "".join(p.capitalize() for p in c.split("_"))) for c in symptoms)],
                        [f"W{counter[0] + 2}", "Illness types a parent can pick",
                         ", ".join(arb[k] for k in ["illnessCold", "illnessFever", "illnessCough", "illnessTummy",
                                                    "illnessEar", "illnessRash", "illnessUnnamed"])],
                        [f"W{counter[0] + 3}", "Footer of the PDF summary for the doctor",
                         "Generated by Palna on [date] from records kept by the family. Not a medical record or "
                         "a diagnosis."]]
            counter[0] += 3
            d.table(["#", "Where it appears", "Exact text shown"], sym_rows, [1.0, 3.4, 7.4],
                    title="Symptom list, illness types and PDF footer")
        if heading.startswith("3.2"):
            # 3.3 growth comes right after medicines.
            d.add("h2", "3.3 Growth check and weight gain")
            d.add("p", "The growth screen plots weight, length and head circumference on the WHO Child Growth "
                       f"Standards (birth to {th['who_max_months']} months) and gives one overall result: "
                       "\u201cTalk to your doctor\u201d, \u201cMention it at the next check-up\u201d or "
                       "\u201cGrowing as expected\u201d (the most serious finding wins). It needs the baby\u2019s "
                       "sex. Only the latest measurement entered after birth is judged; birth values from the "
                       "profile are used only for the birth-weight rule and the weight-gain card. Age is "
                       "chronological (not corrected for prematurity).")
            curves = ", ".join(f"{lbl} ({z:+.2f} SD)" for lbl, z in th["curves"])
            low, high = th["curves"][0], th["curves"][-1]
            n = th["cross_lines"]
            rule_rows = [
                ["G1", "Below the lowest line", f"Latest weight, length or head below the {low[0]} centile "
                 f"({low[1]:.2f} SD).", "Talk to your doctor"],
                ["G2", "Above the highest line (head)", f"Latest head circumference above the {high[0]} centile "
                 f"(+{high[1]:.2f} SD).", "Talk to your doctor"],
                ["G3", "Above the highest line (weight or length)", f"Latest weight or length above the "
                 f"{high[0]} centile.", "Mention it at the next check-up"],
                ["G4", f"Falling across {n} lines", f"Compared with each earlier measurement taken "
                 f"{th['cross_min']} days to {th['cross_max']} days (about 6 months) before the latest one "
                 f"(birth values excluded), the latest value has fallen across {n} or more of the five chart "
                 f"lines: {curves}. The largest crossing found counts.", "Talk to your doctor"],
                ["G5", f"Rising across {n} lines (head)", "As G4, but rising, for head circumference.",
                 "Talk to your doctor"],
                ["G6", f"Rising across {n} lines (weight or length)", "As G4, but rising, for weight or length.",
                 "Mention it at the next check-up"],
                ["G7", "Not back to birth weight", f"A weight measured more than {th['birthweight_days']} days "
                 "after birth is still below the birth weight.", "Talk to your doctor"],
                ["G8", "Slow weight gain", f"Weight gain over about the last month is below the WHO 5th centile "
                 "for 1-month weight increments (birth to 12 months; table below).",
                 "Mention it at the next check-up"],
            ]
            d.table(["#", "Finding", "When the app flags it", "Result shown"], rule_rows,
                    [1.0, 2.9, 6.5, 2.5], title="Growth check rules")
            d.add("p", f"Weight-gain card: the app takes the latest weight and an earlier weight (the birth weight "
                       f"counts) {th['gain_min_days']} to {th['gain_max_days']} days before it, choosing the one "
                       f"closest to one month (30.4 days). The gain is scaled to one month and compared with the "
                       f"WHO 5th, 50th and 95th centiles for the 1-month interval whose middle is closest to the "
                       f"middle of the two dates. Below the 5th centile: \u201cslower than most babies\u201d; above "
                       f"the 95th: \u201cfaster than most babies\u201d. Only up to 12 months.")
            vrows = []
            def interval(x):
                return "Birth" if x == "0w" else x.replace("w", " weeks").replace("m", " months")

            for b, g in zip(velocity["boys"], velocity["girls"]):
                if (b["from"], b["to"]) != (g["from"], g["to"]):
                    raise SourceChanged("WHO weight-gain intervals differ between boys and girls.")
                frm, to = interval(b["from"]), interval(b["to"])
                vrows.append([f"{frm} \u2013 {to}", b["p5"], b["p50"], b["p95"], g["p5"], g["p50"], g["p95"]])
            d.table(["Age interval", "Boys 5th", "Boys 50th", "Boys 95th", "Girls 5th", "Girls 50th",
                     "Girls 95th"], vrows, [3.0, 1.35, 1.35, 1.35, 1.35, 1.35, 1.35],
                    title="WHO weight gain in grams per month (1-month increments), as used by the app")
            wording_table(C.GROWTH_STRINGS, "Growth wording")
        if heading.startswith("3.6"):
            d.add("p", "Hot-weather advice is shown only in: "
                       + ", ".join(C.COUNTRY_NAMES[c] for c in C.COUNTRY_ORDER if c in heat_countries) + ".")
        if heading.startswith("3.7"):
            srows = []
            prev = 0
            for upto, lo, hi in th["sleep"]:
                srows.append([f"{_age_from_days(prev)} \u2013 {_age_from_days(upto)}", f"{lo}\u2013{hi} hours"])
                prev = upto
            d.add("p", "The reports screen shades the typical total sleep per 24 hours for the baby\u2019s age:")
            d.table(["Age", "Typical total sleep per 24 hours"], srows, [5.0, 6.0],
                    title="Typical sleep ranges")
            d.add("p", "For the optional \u201cnext nap\u201d estimate, until a week of the baby\u2019s own sleep "
                       "is logged, the app uses these typical awake times between sleeps:")
            d.table(["Age", "Typical awake time between sleeps"], _awake_rows(th["awake"]), [5.0, 6.0],
                    title="Typical awake times")

    d.sources(C.WORDING_SOURCES)
    d.questions(C.WORDING_QUESTIONS)
    d.add("pagebreak")

    # 4. Emergency numbers -----------------------------------------------------
    d.add("h1", "4. Emergency numbers by country")
    for para in C.EMERGENCY_INTRO:
        d.add("p", para)
    erows = []
    for code in C.COUNTRY_ORDER + [c for c in numbers if c not in C.COUNTRY_ORDER]:
        name = C.COUNTRY_NAMES.get(code, code)
        if code not in C.COUNTRY_ORDER:
            name += " (no vaccine schedule in the app)"
        nums = numbers.get(code)
        if nums is None and code in eu112:
            nums = [("112", "all")]
            name += " (112 only: no country-specific numbers in the app)"
        for j, (num, service) in enumerate(nums or []):
            erows.append([name if j == 0 else "", num, s(SERVICE_LABEL_KEYS[service])])
    erows.append(["Other EU and EEA countries, Switzerland", "112", s("numAll")])
    d.table(["Country", "Number", "Label shown"], erows, [3.6, 2.6, 4.8], title="Emergency numbers")
    d.add("p", f"\u201cCall now\u201d warnings use the first number listed; for Pakistan they say "
               f"\u201c{pk_primary}\u201d.")
    d.sources(C.EMERGENCY_SOURCES)
    d.questions(C.EMERGENCY_QUESTIONS)
    d.add("pagebreak")

    # 5. Solids --------------------------------------------------------------
    d.add("h1", "5. Starting solids: guide, foods and recipes")
    for para in C.SOLIDS_INTRO:
        d.add("p", para)
    g = foods["guides"]
    grows = [["Readiness", g["readiness"]["en"]], ["Start with vegetables", g["vegFirst"]["en"]]]
    grows += [[f"Stage: {st['months'].replace('-', '–')} months", st["text"]["en"]] for st in g["stages"]]
    grows += [["Important rules", r["en"]] for r in g["rules"]]
    grows += [["Allergens", g["allergenAdvice"]["en"]],
              ["Signs of an allergic reaction", fill(g["reactionSigns"]["en"])]]
    grows += [["Myth and fact", f"Belief: {m['belief']['en']}\nWhat we say: {m['evidence']['en']}"]
              for m in g["myths"]]
    d.add("h2", "5.1 Guides")
    d.table(["#", "Topic", "Exact text shown"], [[f"S{n + 1}", t, x] for n, (t, x) in enumerate(grows)],
            [1.1, 2.9, 7.9], title="Solids guides")
    d.table(["#", "Where it appears", "Exact text shown"],
            [[f"S{len(grows) + n + 1}", where, s(k)] for n, (k, where, _) in enumerate(C.SOLIDS_STRINGS)],
            [1.1, 2.9, 7.9], title="Other solids wording")

    d.add("h2", "5.2 Foods")
    allergen_names = {k: v["en"] for k, v in foods["allergens"].items()}
    cat_names = {"fruit": "Fruit", "vegetable": "Vegetable", "grain": "Grain", "legume": "Pulses",
                 "protein": "Meat, fish, egg", "dairy": "Dairy", "nut": "Nuts and seeds", "other": "Other"}
    frows = [[f"F{n + 1}", f"{f['name']['en']}\n({cat_names.get(f['category'], f['category'])})",
              f"{f['fromMonths']} months", f["prep"]["en"],
              allergen_names.get(f["allergen"], "\u2014") if f["allergen"] else "\u2014",
              "Yes" if f["choking"] else "No"] for n, f in enumerate(foods["foods"])]
    d.table(["#", "Food", "From", "Preparation and texture (exact text)", "Allergen", "Choking"], frows,
            [1.0, 2.6, 1.6, 5.0, 1.6, 1.5], title=f"First foods ({len(frows)})")

    d.add("h2", "5.3 Recipes")
    rrows = []
    for n, r in enumerate(foods["recipes"]):
        rrows.append([f"R{n + 1}", f"{r['name']['en']}\nFrom {r['fromMonths']} months",
                      "\n".join("\u2022 " + i["en"] for i in r["ingredients"]),
                      "\n".join(f"{k + 1}. {st['en']}" for k, st in enumerate(r["steps"])),
                      ", ".join(allergen_names[a] for a in r["allergens"]) or "\u2014"])
    d.table(["#", "Recipe and age", "Ingredients", "How to make it", "Allergens"], rrows,
            [1.0, 2.3, 3.7, 4.8, 1.7], title=f"Recipes ({len(rrows)})")
    d.sources(C.SOLIDS_SOURCES)
    d.questions(C.SOLIDS_QUESTIONS)
    d.add("pagebreak")

    # 6. Milestones ----------------------------------------------------------
    d.add("h1", "6. Developmental milestones")
    for para in C.MILESTONES_INTRO:
        d.add("p", para)
    guide = ms["guide"]
    d.table(["#", "Where it appears", "Exact text shown"],
            [["M0a", "Top of the milestones screen", guide["intro"]["en"]],
             ["M0b", "\u201c" + arb["actEarlyTitle"] + "\u201d", guide["actEarly"]["en"]]],
            [1.1, 2.9, 7.9], title="Milestone guidance")
    area_names = {"social": "Social / emotional", "language": "Language", "cognitive": "Thinking",
                  "movement": "Movement"}
    mrows, trows, k, t = [], [], 0, 0
    for age in ms["ages"]:
        for j, m in enumerate(age["milestones"]):
            k += 1
            mrows.append([f"M{k}", f"{age['months']} months" if j == 0 else "",
                          area_names.get(m["area"], m["area"]), m["text"]["en"]])
        for j, tip in enumerate(age["tips"]):
            t += 1
            trows.append([f"T{t}", f"{age['months']} months" if j == 0 else "", tip["en"]])
    d.table(["#", "By age", "Area", "Milestone (most babies, 75% or more)"], mrows,
            [1.0, 1.6, 2.4, 6.8], title=f"Milestones ({len(mrows)})")
    d.table(["#", "Age", "Ways to help (exact text)"], trows, [1.0, 1.6, 9.2],
            title=f"Ways to help ({len(trows)})")
    d.sources([guide["source"].rstrip(".") + ", https://www.cdc.gov/act-early/milestones/ (as cited in the app)."])
    d.questions(C.MILESTONES_QUESTIONS)
    d.add("pagebreak")

    # Final page ---------------------------------------------------------------
    d.add("h1", "General comments")
    d.add("lines", 14)
    d.add("h2", "Sign-off")
    for para in C.SIGN_OFF:
        d.add("p", para)
    d.add("signature", ["Name", "Qualifications / registration no.", "Signature", "Date"])
    d.add("p", f"Please return this pack to {C.CONTACT}. Thank you.")
    return d


def _age_from_days(days: int) -> str:
    if days == 0:
        return "Birth"
    if days % 365 in (0, 1):
        y = days // 365
        return f"{y} year" + ("s" if y != 1 else "")
    return f"{round(days / 30.44)} months"


def _fmt_min(m: int) -> str:
    h, mm = divmod(m, 60)
    if h == 0:
        return f"{mm} min"
    return f"{h} h" + (f" {mm} min" if mm else "")


def _awake_rows(awake) -> list[list[str]]:
    rows, nxt = [], 0
    for pat, lo, hi in awake:
        pat = pat.strip()
        if pat == "_":
            label = f"{nxt} months and older"
        elif pat.startswith("<"):
            upto = int(pat[1:].strip())
            label = f"{nxt}\u2013{upto - 1} months"
            nxt = upto
        else:
            vals = [int(v) for v in pat.split("||")]
            label = ("Under 1 month" if vals == [0] else
                     f"{vals[0]} month" + ("s" if vals[0] != 1 else "") if len(vals) == 1 else
                     f"{vals[0]}\u2013{vals[-1]} months")
            nxt = vals[-1] + 1
        rows.append([label, f"{_fmt_min(int(lo))} \u2013 {_fmt_min(int(hi))}"])
    return rows


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def qid(x: str) -> str:
    """Question ids are shown as PK-Q1, W-Q3 ... so they never clash with row numbers (PK-1, W3)."""
    prefix, num = x.rsplit("-", 1)
    return f"{prefix}-Q{num}"


def md_cell(v) -> str:
    return str(v).replace("|", "\\|").replace("\n", "<br>")


def render_md(d: Doc) -> str:
    out: list[str] = []
    for b in d.blocks:
        kind = b[0]
        if kind == "title":
            out.append(f"# {b[1]}\n")
        elif kind == "subtitle":
            out.append(f"*{b[1]}*\n")
        elif kind == "h1":
            out.append(f"## {b[1]}\n")
        elif kind == "h2":
            out.append(f"### {b[1]}\n")
        elif kind == "h3":
            out.append(f"#### {b[1]}\n")
        elif kind == "p":
            out.append(f"{b[1]}\n")
        elif kind == "bullets":
            out.append("\n".join(f"- {x}" for x in b[1]) + "\n")
        elif kind == "toplist":
            out.append("\n".join(f"{n + 1}. **{qid(q)}** {txt}" for n, (q, txt) in enumerate(b[1])) + "\n")
        elif kind == "questions":
            out.append("\n".join(f"- [ ] **{qid(q)}** {txt}  \n  Answer: ______________________________________"
                                 for q, txt in b[1]) + "\n")
        elif kind == "table":
            headers, rows, _widths, title = b[1], b[2], b[3], b[4]
            if title:
                out.append(f"**{title}**\n")
            lines = ["| " + " | ".join(md_cell(h) for h in headers) + " |",
                     "|" + "|".join("---" for _ in headers) + "|"]
            lines += ["| " + " | ".join(md_cell(c) for c in r) + " |" for r in rows]
            out.append("\n".join(lines) + "\n")
        elif kind == "signature":
            out.append("\n".join(f"**{lbl}:** ________________________________________\n" for lbl in b[1]))
        elif kind == "lines":
            out.append("\n".join("_" * 80 + "\n" for _ in range(b[1])))
        elif kind == "pagebreak":
            out.append("---\n")
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def render_docx(d: Doc, path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.different_first_page_header_footer = True

    navy = RGBColor(0x1F, 0x38, 0x64)

    def set_font(style, size, bold=None, color=None):
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Calibri")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]

    styles = doc.styles
    set_font(styles["Normal"], 11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    set_font(styles["Title"], 26, True, navy)
    set_font(styles["Heading 1"], 17, True, navy)
    set_font(styles["Heading 2"], 13.5, True, navy)
    set_font(styles["Heading 3"], 11.5, True, navy)
    set_font(styles["List Bullet"], 11)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.space_before = Pt(12 if name != "Heading 3" else 8)

    # Footer: "Palna: medical content review · 28 October 2026 · Page X of Y"
    def add_field(paragraph, instr):
        run = paragraph.add_run()
        for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
                run._r.append(el)
            elif text == instr:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {instr} "
                run._r.append(el)
            else:
                el = OxmlElement("w:t")
                el.text = text
                run._r.append(el)
        run.font.size = Pt(9)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{C.TITLE} \u00b7 {C.VERSION_DATE} \u00b7 Page ")
    r.font.size = Pt(9)
    add_field(fp, "PAGE")
    r = fp.add_run(" of ")
    r.font.size = Pt(9)
    add_field(fp, "NUMPAGES")

    def shade(cell, fill_hex):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), fill_hex)
        tcpr.append(sh)

    def row_props(row, header=False):
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trpr.append(el)
        if header:
            el = OxmlElement("w:tblHeader")
            trpr.append(el)

    def cell_text(cell, text, bold=False, size=9.5):
        cell.text = ""
        lines = str(text).split("\n")
        p = cell.paragraphs[0]
        for k, line in enumerate(lines):
            if k:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.size = Pt(size)
            run.bold = bold

    for b in d.blocks:
        kind = b[0]
        if kind == "title":
            for _ in range(5):
                doc.add_paragraph()
            doc.add_paragraph(b[1], style="Title")
        elif kind == "subtitle":
            p = doc.add_paragraph()
            run = p.add_run(b[1])
            run.font.size = Pt(13)
            run.italic = True
            run.font.color.rgb = navy
        elif kind in ("h1", "h2", "h3"):
            doc.add_heading(b[1], level={"h1": 1, "h2": 2, "h3": 3}[kind])
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "bullets":
            for x in b[1]:
                doc.add_paragraph(x, style="List Bullet")
        elif kind == "toplist":
            for n, (q, txt) in enumerate(b[1]):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.8)
                p.add_run(f"{n + 1}.\t")
                p.add_run(qid(q) + "  ").bold = True
                p.add_run(txt)
        elif kind == "questions":
            for q, txt in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.first_line_indent = Cm(-0.7)
                p.add_run("\u2610\t")
                p.add_run(qid(q) + "  ").bold = True
                p.add_run(txt)
                a = doc.add_paragraph()
                a.paragraph_format.left_indent = Cm(0.7)
                a.paragraph_format.space_after = Pt(8)
                run = a.add_run("Answer: " + "_" * 70)
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif kind == "table":
            headers, rows, widths, title = b[1], b[2], b[3], b[4]
            if title:
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(title)
                run.bold = True
                run.font.size = Pt(10.5)
            if len(widths) != len(headers):
                raise ValueError(f"Table {title!r}: {len(widths)} widths for {len(headers)} columns")
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            tblpr = t._tbl.tblPr
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tblpr.append(layout)
            for i, h in enumerate(headers):
                c = t.rows[0].cells[i]
                cell_text(c, h, bold=True, size=9)
                shade(c, "D9E2F3")
            row_props(t.rows[0], header=True)
            for ri, row in enumerate(rows):
                tr = t.rows[ri + 1]
                row_props(tr)
                for ci, v in enumerate(row):
                    cell_text(tr.cells[ci], v)
                    if ci >= len(headers) - 2:
                        shade(tr.cells[ci], "FFF9E6")
            for row in t.rows:
                for ci, w in enumerate(widths):
                    row.cells[ci].width = Cm(w)
            grid = t._tbl.tblGrid
            for ci, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(Cm(widths[ci]).twips)))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "signature":
            for lbl in b[1]:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.add_run(f"{lbl}:").bold = True
                p.add_run("\t" + "_" * 55)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(6.0))
        elif kind == "lines":
            for _ in range(b[1]):
                p = doc.add_paragraph("_" * 86)
                p.paragraph_format.space_before = Pt(10)
                p.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.core_properties.title = C.TITLE
    doc.core_properties.subject = "Medical content review for a pediatrician"
    doc.core_properties.author = "Palna"
    doc.save(path)


def main() -> int:
    try:
        d = build()
    except SourceChanged as e:
        print("STOPPED:", e, file=sys.stderr)
        return 1
    OUT_MD.write_text(render_md(d), encoding="utf-8")
    render_docx(d, OUT_DOCX)
    print(f"Wrote {OUT_MD.name} and {OUT_DOCX.name}: {d.table_count} tables, {d.question_count} open questions.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
